"""Portable text blocks for Word/PDF. No HTML execution or remote image fetching."""
import re


def blocks(value):
    result=[];paragraph=[];code=None
    def flush():
        if paragraph:
            text=' '.join(paragraph);chunk=''
            for sentence in re.split(r'(?<=[.!?])\s+',text):
                if len(chunk)>400:
                    result.append(('paragraph',chunk.strip()));chunk=''
                chunk+=sentence+' '
            if chunk.strip():result.append(('paragraph',chunk.strip()))
            paragraph.clear()
    for raw in str(value or '').replace('\r','').split('\n'):
        line=raw.strip()
        if line.startswith('```'):
            flush()
            if code is not None:result.append(('code','\n'.join(code)));code=None
            else:code=[]
            continue
        if code is not None:code.append(raw);continue
        if not line:flush();continue
        heading=re.match(r'^#{1,3}\s+(.+)$',line)
        bullet=re.match(r'^(?:[-*•]|\d+[.)])\s+(.+)$',line)
        if heading:flush();result.append(('heading',heading[1]))
        elif bullet:flush();result.append(('bullet',bullet[1]))
        elif line.startswith('> '):flush();result.append(('callout',line[2:]))
        else:paragraph.append(line)
    flush()
    if code is not None:result.append(('code','\n'.join(code)))
    return result


def add_docx(doc,value,heading_level=3):
    from docx.shared import Pt,RGBColor
    for kind,text in blocks(value):
        if kind=='heading':doc.add_heading(text,heading_level)
        elif kind=='bullet':doc.add_paragraph(text,'List Bullet')
        else:
            p=doc.add_paragraph(text)
            p.paragraph_format.space_after=Pt(10)
            p.paragraph_format.line_spacing=1.35
            if kind=='code':
                for run in p.runs:run.font.name='Consolas';run.font.size=Pt(10)
            elif kind=='callout':
                for run in p.runs:run.italic=True;run.font.color.rgb=RGBColor.from_string('6944B7')


def lecture_issues(lecture,previous=()):
    """Reject plans/promises and near-copies; factual review still belongs to people."""
    text=str(lecture.get('content') or '')
    issues=[]
    if len(text.strip())<1600:issues.append('Нужен самостоятельный учебный текст от 1600 символов, а не краткий план')
    sections=[];current=None
    for kind,value in blocks(text):
        if kind=='heading':
            current=[value.lower(),[]];sections.append(current)
        elif current is not None:current[1].append(value)
    def section(pattern):
        return '\n'.join('\n'.join(parts) for title,parts in sections if re.search(pattern,title))
    explanation=section(r'объясн|поняти|как работает|теори|основ|разбор темы')
    example=section(r'пример|кейс|разбор задач')
    practice=section(r'попроб|практик|упражн|задани')
    check=section(r'самопроверк|вопрос.*ответ|провер.*поним')
    if len(explanation)<450:issues.append('Добавь раздел «Объяснение понятий»: определения и объяснение механизма, от 450 символов')
    if len(example)<220 or not re.search(r'результат|получ|ответ|вывод|итог',example,re.I):
        issues.append('Добавь «Разобранный пример»: исходные данные, конкретные шаги и полученный результат, от 220 символов')
    if len(practice)<100:issues.append('Добавь «Попробуйте сами»: конкретное упражнение с исходными данными, от 100 символов')
    if len(check)<120 or not re.search(r'ответ|решени|ключ',check,re.I):issues.append('Добавь «Самопроверка»: вопросы и содержательные ответы, от 120 символов')
    sentences=[x.strip() for x in re.split(r'[.!?]\s+',text) if len(x.strip())>30]
    promises=sum(bool(re.search(r'мы (?:будем|рассмотрим|изучим|познакомимся|узнаем|научимся|разбер[её]м)|вы (?:узнаете|научитесь|изучите)',x,re.I)) for x in sentences)
    if promises>=3 and promises/max(1,len(sentences))>.35:issues.append('Замени обещания «мы изучим/вы научитесь» фактическим объяснением темы и решениями')
    def shingles(value):
        words=re.findall(r'[а-яa-z0-9]+',value.lower())
        return {tuple(words[i:i+5]) for i in range(max(0,len(words)-4))}
    own=shingles(text)
    if len(own)>100:
        for other in previous:
            if other.get('topicId')==lecture.get('topicId'):continue
            prior=shingles(str(other.get('content') or ''))
            if len(prior)>100 and len(own&prior)/min(len(own),len(prior))>.88:
                issues.append('Текст почти повторяет другой урок: '+str(other.get('title',''))[:180]+'. Раскрой именно текущую тему новым объяснением и примером')
                break
    return issues


def normalize_lecture(value):
    if not isinstance(value,dict):return value
    if 'explanation' in value:
        from .contracts import GeneratedLecture
        lesson=GeneratedLecture.model_validate(value).model_dump()
        material={key:lesson[key] for key in ('id','topicId','competencyIds','title','plan','slides')}
        result=[]
        def add(kind,text):result.append({'kind':kind,'text':text,'lines':[]})
        add('heading','Объяснение понятий')
        for paragraph in lesson['explanation']:add('paragraph',paragraph)
        add('heading','Разобранный пример');example=lesson['workedExample']
        add('paragraph',example['context'])
        for step in example['steps']:add('bullet',step)
        for code in example['code']:result.append({'kind':'code','text':code['language'],'lines':code['lines']})
        add('callout','Результат: '+example['result'])
        add('heading','Попробуйте сами');add('paragraph',lesson['practice']['task'])
        add('callout','Ожидаемый результат: '+lesson['practice']['expectedResult'])
        add('heading','Типичные ошибки и исправления')
        for mistake in lesson['mistakes']:
            add('bullet',mistake['error']);add('paragraph',mistake['correction'])
        add('heading','Самопроверка с ответами')
        for check in lesson['selfCheck']:
            add('paragraph','Вопрос: '+check['question']);add('paragraph','Ответ: '+check['answer'])
        add('heading','Краткие выводы')
        for takeaway in lesson['takeaways']:add('bullet',takeaway)
        return normalize_lecture({**material,'blocks':result})
    if 'blocks' not in value:return value
    from .contracts import BlockLecture
    import ast
    lesson=BlockLecture.model_validate(value).model_dump()
    parts=[]
    for block in lesson.pop('blocks'):
        kind,text,lines=block['kind'],block['text'].strip(),block['lines']
        if kind=='heading':parts.append('## '+text)
        elif kind=='bullet':parts.append('- '+text)
        elif kind=='callout':parts.append('> '+text)
        elif kind=='code':
            if not lines:raise ValueError('Для блока code нужны отдельные строки в lines с сохранёнными отступами')
            language=text.lower() if re.fullmatch(r'[a-zA-Z0-9+#.-]{0,30}',text) else ''
            code='\n'.join(line.replace('\r\n','\n').replace('\r','\n') for line in lines)
            if language in ('py','python'):
                try:ast.parse(code)
                except SyntaxError as error:raise ValueError('Исправь синтаксис Python и отступы в блоке code: '+str(error)) from None
            parts.append('```'+language+'\n'+code+'\n```')
        else:parts.append(text)
    lesson['content']='\n\n'.join(parts)
    for slide in lesson.get('slides',[]):
        if slide.get('layout')=='steps':slide['text']=re.sub(r'(?<!\d)\s+(?=\d+[.)]\s)', '\n',slide['text'])
    if len(lesson['content'])>60000:raise ValueError('Учебный текст превышает 60000 символов')
    return lesson
