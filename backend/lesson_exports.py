from io import BytesIO
import html
from fastapi.responses import Response
from docx import Document
from docx.shared import Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Preformatted
import os

from .platform import lessons
from .programs_helpers import ensure
from .slides import program_pptx
from .learning_content import blocks, add_docx


def export_lesson(program,lesson_id,format):
    lesson = next((l for l in lessons(program) if l["id"] == lesson_id),None)
    ensure(lesson,"Урок не найден",404)
    output=BytesIO()
    if format == "docx":
        doc=Document()
        doc.styles['Normal'].font.name='Arial'
        doc.styles['Normal'].font.size=Pt(12)
        doc.add_heading(lesson['title'],0)
        doc.add_paragraph('Эдентика · Черновик для экспертной проверки')
        add_docx(doc,lesson['text'],heading_level=2)
        doc.save(output)
        mime='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    elif format == "pptx":
        single={"skeleton":program['skeleton'],"materials":{"one":{"lectures":[{"slides":lesson['slides']}]}}}
        output.write(program_pptx(single))
        mime='application/vnd.openxmlformats-officedocument.presentationml.presentation'
    elif format == "pdf":
        font=os.getenv('CERTIFICATE_FONT') or ('C:/Windows/Fonts/arial.ttf' if os.name=='nt' else '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')
        if 'Edentica' not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont('Edentica',font))
        styles=getSampleStyleSheet()
        styles.add(ParagraphStyle(name='BodyRu',fontName='Edentica',fontSize=12,leading=20,spaceAfter=12))
        styles.add(ParagraphStyle(name='CodeRu',fontName='Edentica',fontSize=10,leading=15,spaceAfter=12,backColor='#F3F0FA',leftIndent=12))
        styles.add(ParagraphStyle(name='SectionRu',fontName='Edentica',fontSize=16,leading=22,spaceBefore=14,spaceAfter=10,keepWithNext=True))
        styles.add(ParagraphStyle(name='TitleRu',fontName='Edentica',fontSize=22,leading=30,spaceAfter=20))
        story=[]
        story.append(Paragraph(html.escape(lesson['title']),styles['TitleRu']))
        for kind,text in blocks(lesson['text']):
            if kind=='code':
                story.append(Preformatted(text,styles['CodeRu'],maxLineLength=90))
                continue
            style=styles['SectionRu' if kind=='heading' else 'BodyRu']
            safe=html.escape(text).replace('\n','<br/>')
            if kind=='bullet':safe='• '+safe
            story.append(Paragraph(safe,style))
        SimpleDocTemplate(output,pagesize=A4).build(story)
        mime='application/pdf'
    else:
        ensure(False,"Неподдерживаемый формат")
    extension=format
    return Response(output.getvalue(),media_type=mime,headers={'Content-Disposition':f'attachment; filename="lesson.{extension}"'})
