"""Portable editable PPTX export without host-specific Node packages."""
from io import BytesIO

from fastapi import HTTPException
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
import re


def program_pptx(program, module_id=None, topic_ids=None):
    materials = program.get("materials", {})
    if module_id and module_id not in materials:
        raise HTTPException(409, "Материалы модуля ещё не готовы")
    modules = [materials[module_id]] if module_id else list(materials.values())
    slides = [s for m in modules for lecture in m.get("lectures", []) if topic_ids is None or lecture.get('topicId') in topic_ids for s in lecture.get("slides", [])]
    if not 1 <= len(slides) <= 300:
        raise HTTPException(422, "Выгружайте отдельный модуль: допустимо 1–300 слайдов")
    skeleton = program["skeleton"]
    slides.insert(0, {"title": skeleton["title"], "text": "Черновик для экспертной проверки. " + skeleton.get("audience", ""),
                      "notes": "Содержание требует утверждения человеком."})
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    for index, content in enumerate(slides):
        # Split lengthy model text into additional slides instead of clipping it.
        text = str(content.get("text", ""))
        chunks = [text[i:i + 1000] for i in range(0, len(text), 1000)] or [""]
        for part, chunk in enumerate(chunks):
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = RGBColor.from_string("F7F8FA")

            def box(value, x, y, width, height, size, color, bold=False):
                frame = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height)).text_frame
                frame.word_wrap = True
                frame.text = str(value)
                for paragraph in frame.paragraphs:
                    paragraph.font.name = "DejaVu Sans"
                    paragraph.font.size = Pt(size)
                    paragraph.font.bold = bold
                    paragraph.font.color.rgb = RGBColor.from_string(color)
                return frame

            box("ЭДЕНТИКА · ЧЕРНОВИК", .6, .3, 12, .4, 12, "7A52E9")
            title = str(content.get("title", ""))
            if part:
                title += " · продолжение"
            box(title, .6, 1, 12, 1.5, 20 if len(title)>150 else 26 if len(title) > 90 else 32, "202126", True)
            layout=content.get('layout') or ('concept','steps','example','summary')[index%4]
            def card(x,y,w,h,color='EEE9FA'):
                shape=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
                shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string(color)
                shape.line.fill.background()
            labels={'concept':'КЛЮЧЕВОЙ ПРИНЦИП','steps':'ШАГ ЗА ШАГОМ','comparison':'СРАВНЕНИЕ','example':'РАЗБОР ПРИМЕРА','summary':'ЧТО ЗАПОМНИТЬ'}
            box(labels.get(layout,'МАТЕРИАЛ УРОКА'),.65,2.25,11,.3,11,'7A52E9',True)
            rows=[re.sub(r'^\s*(?:[-*•]|\d+[.)])\s+','',x).strip() for x in chunk.split('\n') if x.strip()]
            if layout=='steps' and 2<=len(rows)<=5:
                height=min(1.0,3.9/len(rows))
                for n,row in enumerate(rows):
                    y=2.8+n*height;card(.6,y,.55,.5,'C5F779');box(str(n+1),.72,y+.05,.35,.35,16,'202126',True)
                    box(row,1.4,y,11,height-.07,17 if len(row)>130 else 20,'202126')
            elif layout=='comparison' and len(rows)>=2:
                # Keep both alternatives visible, without discarding any rows.
                middle=(len(rows)+1)//2
                for n,part_rows in enumerate((rows[:middle],rows[middle:])):
                    x=.6+6.2*n;card(x,2.8,5.9,3.8,'EEE9FA' if n==0 else 'EAF5DE')
                    box('\n\n'.join(part_rows),x+.25,3.05,5.4,3.3,16 if len(chunk)>600 else 20,'202126')
            elif layout=='example':
                card(.6,2.8,12,3.85,'24242D')
                box(chunk,.9,3.05,11.4,3.3,16 if len(chunk)>600 else 20,'FFFFFF')
            elif layout=='summary':
                card(.6,2.8,12,3.85,'EAF5DE')
                box(chunk,.9,3.05,11.4,3.3,16 if len(chunk)>600 else 20,'202126')
            else:
                box(chunk,.6,2.8,10.7,3.85,16 if len(chunk)>600 else 20,'202126')
                card(12,2.8,.3,3.7,'C5F779')
            box(str(len(deck.slides)), 12, 7, .7, .3, 11, "727780")
            slide.notes_slide.notes_text_frame.text = str(content.get("notes", ""))
    output = BytesIO()
    deck.save(output)
    return output.getvalue()
