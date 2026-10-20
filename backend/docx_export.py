"""Build a draft using retained ACR reference section geometry and styles."""
import io
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent

def build(p):
    from .canonical import canonicalize_sections
    p={**p,'sections':canonicalize_sections(p.get('sections',{}),p['skeleton'])}
    s = p['skeleton']
    template = ROOT / 'templates' / ('acr-dpp.docx' if s['type'] == 'dpp' else 'acr-general.docx')
    doc = Document(template)
    # Retain styles, relationships, section dimensions and page furniture from the reference.
    for child in list(doc.element.body):
        if child.tag != qn('w:sectPr'):
            doc.element.body.remove(child)
    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(12)
    normal.paragraph_format.space_after = Pt(6)
    for name in ['Title', 'Heading 1', 'Heading 2', 'Heading 3']:
        doc.styles[name].font.name = 'Times New Roman'
        doc.styles[name].font.color.rgb = __import__('docx').shared.RGBColor(0, 0, 0)
    for section in doc.sections:
        # The authoritative reference uses A4. Keep it; constrain tables to printable width.
        section.header_distance = Cm(1)
        section.footer_distance = Cm(1)
    def para(value, bold=False):
        paragraph = doc.add_paragraph()
        paragraph.add_run(str(value)).bold = bold
        return paragraph
    def heading(value, level=1):
        h = doc.add_heading(value, level)
        h.paragraph_format.keep_with_next = True
        return h
    def table(headers, rows):
        t = doc.add_table(rows=1, cols=len(headers))
        t.autofit = False
        width = doc.sections[-1].page_width - doc.sections[-1].left_margin - doc.sections[-1].right_margin
        weights = [max(1, len(x)) for x in headers]
        if len(headers) >= 5:
            weights = [40] + [10] * (len(headers)-1)
        for i, h in enumerate(headers):
            t.columns[i].width = int(width*weights[i]/sum(weights))
            t.rows[0].cells[i].text = h
        repeat = OxmlElement('w:tblHeader'); repeat.set(qn('w:val'), 'true'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for values in rows:
            cells = t.add_row().cells
            for c, value in zip(cells, values): c.text = str(value)
        for ri, row in enumerate(t.rows):
            for cell in row.cells:
                props = cell._tc.get_or_add_tcPr()
                borders = OxmlElement('w:tcBorders')
                for edge in ['top', 'left', 'bottom', 'right']:
                    b = OxmlElement('w:'+edge); b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), '4'); b.set(qn('w:color'), 'D9D9D9'); borders.append(b)
                props.append(borders)
                if ri == 0:
                    fill = OxmlElement('w:shd'); fill.set(qn('w:fill'), 'EDEDED'); props.append(fill)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(4)
                    for run in paragraph.runs: run.font.size = Pt(10); run.bold = ri == 0
        return t
    for value in ['Академия цифрового развития', 'Организация и реквизиты ____________________', 'ЧЕРНОВИК ДЛЯ СОГЛАСОВАНИЯ', 'УТВЕРЖДЕНО', 'Приказ № ______ от «___» __________ 20___ г.']:
        para(value).alignment = WD_ALIGN_PARAGRAPH.CENTER
    kind = 'Дополнительная профессиональная программа повышения квалификации' if s['type']=='dpp' else 'Дополнительная общеобразовательная общеразвивающая программа'
    para(kind, True).alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(s['title'], 'Title')
    for value in [str(s['totalHours'])+' академических часов', 'Форма обучения: '+s['format'], 'Преподаватели ____________________', 'Подписи ____________________', 'Магас '+str(__import__('datetime').date.today().year)]:
        para(value).alignment = WD_ALIGN_PARAGRAPH.CENTER
    para('Сформировано агентом Edentica. Требует проверки методистом и экспертом. Неподтверждённые источники указаны в приложении.')
    if p.get('demo'): para('ДЕМОНСТРАЦИОННЫЙ ПРИМЕР. НЕ РЕЗУЛЬТАТ ГЕНЕРАЦИИ МОДЕЛЬЮ.', True)
    doc.add_page_break()
    headings = p.get('templateSections', list(p['sections']))
    for title in headings:
        heading(title)
        if title!='Планируемые результаты обучения':para(p['sections'].get(title, 'Требует заполнения методистом'))
        if title == 'Планируемые результаты обучения':
            for c in s['competencies']:
                heading(c['id']+' '+c['title'], 2)
                para('Знать: '+c['knowledge']);para('Уметь: '+c['abilities']);para('Владеть: '+c['skills'])
                if c.get('standard'):
                    standard=c['standard']
                    para('Профстандарт: '+standard.get('title','')+'; код '+standard.get('code','')+'. Трудовая функция '+standard.get('functionCode','')+' — '+standard.get('functionTitle','требует проверки')+'.')
                    if standard.get('mappingReason'):para('Связь с результатом обучения: '+standard['mappingReason'])
                    para('Источник: '+standard.get('url','не подтверждён')+'. Актуальность и применимость подтверждает эксперт.')
        if title in ['Учебный план', 'Учебно-тематический план']:
            rows=[]
            for m in s['modules']:
                if title=='Учебный план':
                    hs={k:sum(t['hours'][k] for t in m['topics']) for k in ['lecture','practice','lab','self']}
                    rows.append([m['title'],m['totalHours'],*hs.values(),0])
                else:
                    rows.extend([[t['id']+' '+t['title'],t['totalHours'],*[t['hours'][k] for k in ['lecture','practice','lab','self']],0] for t in m['topics']])
            rows.append(['Итоговая аттестация',s['finalHours'],0,0,0,0,s['finalHours']])
            totals=[sum(t['hours'][k] for m in s['modules'] for t in m['topics']) for k in ['lecture','practice','lab','self']]
            rows.append(['Итого',s['totalHours'],*totals,s['finalHours']])
            table(['Модуль или тема','Всего','Лекции','Практики','Лаб.','СРС','Аттест.'],rows)
            if title=='Учебно-тематический план' and any(t.get('delivery',{}).get('remote') for m in s['modules'] for t in m['topics']):
                para('Предложенное распределение контактных часов по форме участия. Соотношение утверждает методист. Часы самостоятельной работы и итоговой аттестации показаны отдельно в основном плане.')
                deliveries=[[t['id']+' '+t['title'],t.get('delivery',{}).get('inPerson',0)+t.get('delivery',{}).get('remote',0),t.get('delivery',{}).get('inPerson',0),t.get('delivery',{}).get('remote',0)] for m in s['modules'] for t in m['topics']]
                deliveries.append(['Итого',*[sum(row[i] for row in deliveries) for i in range(1,4)]])
                table(['Тема','Контактные часы','Очно','Дистанционно'],deliveries)
        if title=='Календарный учебный график':
            table(['Занятие','Неделя / день','Часы','Темы и часы'],[[i+1,str(x['week'])+' / '+str(x.get('day','—')),x['hours'],
                '; '.join(('Итоговая аттестация' if a['topicId']=='final' else a['topicId'])+' — '+str(a['hours'])+' ч.' for a in x.get('allocations',[{'topicId':x['topicIds'][0],'hours':x['hours']}]))] for i,x in enumerate(s['calendar']['sessions'])])
        if title in ['Рабочая программа учебных дисциплин','Содержание программы']:
            for m in s['modules']:
                heading(m['title'],2)
                content=p['materials'].get(m['id'],{})
                para(content.get('summary',''));para(content.get('methodology',''))
                for topic in m['topics']: para(topic['id']+' '+topic['title']+' — '+str(topic['totalHours'])+' ч. Результаты: '+', '.join(topic['competencyIds']))
    heading('Приложение 1 Оценочные материалы')
    for m in p['materials'].values():
        heading(next(x['title'] for x in s['modules'] if x['id']==m['moduleId']),2)
        for i,q in enumerate(m['questions'],1):
            para(str(i)+'. '+q['text'],True)
            for j,option in enumerate(q['options'],1):para(str(j)+') '+option)
            para('Ключ: '+str(q['answer']+1)+'. '+q['explanation'])
            para('Результаты: '+', '.join(q['competencyIds'])+'; тема '+q['topicId'])
        for a in m['seminars']:
            heading(a['title'],3);para(a['task']);para('Сдаваемый результат: '+a['deliverable'])
            table(['Критерий','Баллы'],[[c['title'],c['points']] for c in a['criteria']])
        for a in m.get('interviews',[]):
            heading('Собеседование: '+a['topicId'],3);para(a['text']);para('Ожидаемый ответ: '+a['expectedAnswer'])
            para('Результаты: '+', '.join(a['competencyIds']))
            table(['Критерий','Баллы'],[[c['title'],c['points']] for c in a['criteria']])
    f=p.get('finalAssessment')
    if f:
        heading('Итоговая аттестация',2);para(f['title']);para(f['task']);para(f.get('deliverable',''))
        for item in f.get('items',[]):para(item['title']+': '+item['requirements'])
        table(['Критерий','Баллы'],[[c['title'],c['points']] for c in f['criteria']])
    heading('Приложение 2 Руководства самостоятельного изучения и слайды' if s.get('selfStudyOnly') else 'Приложение 2 Планы лекций')
    for m in p['materials'].values():
        for lecture in m['lectures']:
            heading(lecture['title'],2);para(lecture['plan']);para(lecture.get('content',''))
            for i,slide in enumerate(lecture['slides'],1):
                heading('Слайд '+str(i)+' '+slide['title'],3);para(slide['text']);para('Заметки преподавателя: '+slide['notes'])
    heading('Приложение 3 Источники и отчёт проверки')
    for source in p.get('sources',[]):
        verification=source.get('verification',{})
        state={'confirmed':'ПОДТВЕРЖДЁН МЕТОДИСТОМ','rejected':'ОТКЛОНЁН'}.get(verification.get('status'),'НЕ ПОДТВЕРЖДЁН')
        para(source['title']+' — '+source.get('url','')+' ['+state+']')
        if verification.get('evidence'):para('Основание: '+verification['evidence'])
    report=p.get('report',{})
    para('Ошибок: '+str(len(report.get('errors',[])))+'. Замечаний: '+str(len(report.get('warnings',[]))))
    for issue in report.get('errors',[])+report.get('warnings',[]):para(issue['message'])
    coverage=report.get('coverage',[])
    if coverage:table(['Компетенция','Темы','Вопросы','Задания'],[[x['id'],', '.join(x['topics']),', '.join(x['questions']),', '.join(x['assignments'])] for x in coverage])
    for section in doc.sections:
        footer=section.footer.paragraphs[0]
        footer.text='Черновик Edentica · Не утверждён · '
        field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    buffer=io.BytesIO();doc.save(buffer);return buffer.getvalue()

