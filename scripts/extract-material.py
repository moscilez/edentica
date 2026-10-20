import io, json, sys, zipfile
import os

# Keep hostile parsers within a child-process budget on Linux servers.
if os.name != 'nt':
    import resource
    resource.setrlimit(resource.RLIMIT_AS,(384*1024*1024,384*1024*1024))
    resource.setrlimit(resource.RLIMIT_CPU,(15,15))

kind = sys.argv[1]
data = sys.stdin.buffer.read()
if len(data) > 5_000_000: raise ValueError('File too large')
if kind == 'txt':
    result = data.decode('utf-8-sig')
elif kind == 'pdf':
    from pypdf import PdfReader
    if not data.startswith(b'%PDF-'): raise ValueError('Invalid PDF')
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted or len(reader.pages) > 100: raise ValueError('Unsupported PDF')
    parts=[]
    extracted=0
    for page in reader.pages:
        text=page.extract_text() or ''
        parts.append(text[:51000-extracted])
        extracted+=len(parts[-1])
        if extracted>50000:break
    result = '\n\n'.join(parts)
else:
    if not data.startswith(b'PK'): raise ValueError('Invalid Office file')
    archive = zipfile.ZipFile(io.BytesIO(data))
    entries = archive.infolist()
    if len(entries) > 1000 or sum(e.file_size for e in entries) > 20_000_000: raise ValueError('Archive limits exceeded')
    for entry in entries:
        if entry.filename.endswith('.xml'):
            content = archive.read(entry)
            markers = content.replace(b'\x00', b'').upper()
            if b'<!DOCTYPE' in markers or b'<!ENTITY' in markers: raise ValueError('External entities forbidden')
    if kind == 'docx':
        from docx import Document
        doc = Document(io.BytesIO(data))
        result = '\n\n'.join([p.text for p in doc.paragraphs] + [' | '.join(cell.text for cell in row.cells) for table in doc.tables for row in table.rows])
    elif kind == 'pptx':
        from pptx import Presentation
        deck = Presentation(io.BytesIO(data))
        if len(deck.slides) > 100: raise ValueError('Too many slides')
        result = '\n\n'.join(shape.text for slide in deck.slides for shape in slide.shapes if hasattr(shape, 'text'))
    else:
        raise ValueError('Unsupported format')
result = result.strip()
if not result: raise ValueError('No text found')
sys.stdout.buffer.write(json.dumps({'text':result[:50000],'truncated':len(result)>50000},ensure_ascii=False).encode('utf-8'))
