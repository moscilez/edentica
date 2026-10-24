import asyncio
import copy
from io import BytesIO
from uuid import uuid4

import pytest
from docx import Document
from fastapi import HTTPException
from fastapi.testclient import TestClient
from pptx import Presentation

from backend.app import create_app
from backend.lesson_planning import lesson_slots,coverage
from backend.lesson_exports import export_lesson
from backend.platform import lessons
from backend.validation import inspect_module
from tests_python.test_workflow import fixture,semantic


def test_ten_lessons_have_independent_materials_and_three_exports():
    p=fixture();module=copy.deepcopy(p['skeleton']['modules'][0])
    module['topics']=module['topics'][:1]
    module['topics'][0].update(totalHours=20,lessonCount=10,hours={'lecture':10,'practice':10,'lab':0,'self':0})
    module['totalHours']=20
    module['lessonPlan']=lesson_slots(module)
    assert len(module['lessonPlan'])==10
    assert sum(x['hours'] for x in module['lessonPlan'])==20
    assert all(sum(x['activities'].values())==x['hours'] for x in module['lessonPlan'])
    assert sum(x['activities']['lecture'] for x in module['lessonPlan'])==10
    material=copy.deepcopy(next(iter(p['materials'].values())))
    original=material['lectures'][0]
    material['moduleId']=module['id'];material['lectures']=[];material['questions']=[]
    for n,slot in enumerate(module['lessonPlan'],1):
        material['lectures'].append({**original,'id':slot['id'],'lessonId':slot['id'],
            'topicId':slot['topicId'],'title':f'Отдельная тема {n}','content':f'Материал урока {n}. '+original['content']})
        material['questions'].append({**next(iter(p['materials'].values()))['questions'][0],
            'id':f'Q{n}','lessonId':slot['id'],'topicId':slot['topicId']})
    p['materials']={module['id']:material}
    assert coverage(module,material)['completed']==10
    course_lessons=lessons(p)
    assert len(course_lessons)==10
    assert all(len(l['questions'])==1 for l in course_lessons)
    assert len({l['title'] for l in course_lessons})==10
    for n,lesson in enumerate(course_lessons,1):
        for fmt in ('docx','pptx','pdf'):
            response=export_lesson(p,lesson['id'],fmt)
            assert response.body
            if fmt=='docx':assert f'Материал урока {n}' in '\n'.join(x.text for x in Document(BytesIO(response.body)).paragraphs)
            if fmt=='pptx':assert len(Presentation(BytesIO(response.body)).slides)>1
            if fmt=='pdf':assert response.body.startswith(b'%PDF')
    material['lectures'].pop()
    assert coverage(module,material)['completed']==9
    assert any('отдельные уроки' in e['message'] for e in inspect_module(material,module,p['skeleton']))


@pytest.mark.parametrize('fail',[False,True])
def test_brief_revision_rebuilds_whole_structure_atomically(postgres_url,fail):
    calls=[]
    async def model(skill,payload):
        calls.append(payload)
        if fail:raise HTTPException(503,'Тестовая недоступность модели')
        outline=semantic()
        outline['title']='Кибербезопасность для студентов'
        outline['modules'][0]['title']='Защита учётных записей'
        outline['modules'][0]['topics'][0]['title']='Кибербезопасность: защита паролей'
        outline['competencies'][0]['title']='Настраивать защиту аккаунта'
        return {'value':outline,'usage':{'inputTokens':1,'outputTokens':1},'model':'test'}
    app=create_app(postgres_url,model)
    with TestClient(app,base_url='http://localhost') as client:
        app.state.programs.allow_unconfigured=True
        h={'X-Edentica':'1'}
        client.post('/api/auth/register',headers=h,json={'login':'brief-'+uuid4().hex,'name':'Тестовый методист','password':'Disposable-only-123','role':'methodist'}).raise_for_status()
        def post(action,**body):
            r=client.post('/api/agent',headers=h,json={'action':action,**body});assert r.status_code in (200,201),r.text
            return r.json()
        before=post('demo')
        brief={k:before['skeleton'][k] for k in ('title','goal','audience','entryLevel','format','assessment')}
        brief.update(goal='Научить студентов защищать свои аккаунты',audience='Студенты колледжа',totalHours=24,weeks=4,requiredTopics=['кибербезопасность'])
        post('revise_brief',id=before['id'],brief=brief,expectedRevision=before['revision'])
        import time
        deadline=time.monotonic()+10
        while time.monotonic()<deadline:
            after=post('get',id=before['id'])
            if after['job']['state']!='running':break
            time.sleep(.02)
        assert calls
        if fail:
            assert after['job']['state']=='failed'
            for key in ('skeleton','materials','requirements','approval','revision'):
                assert after[key]==before[key]
        else:
            assert after['job']['state']=='completed',after['job']
            assert after['skeleton']['goal']==brief['goal']
            assert after['skeleton']['title']=='Кибербезопасность для студентов'
            assert after['skeleton']['modules'][0]['title']=='Защита учётных записей'
            assert after['skeleton']['competencies'][0]['title']=='Настраивать защиту аккаунта'
            assert after['skeleton']['totalHours']==24
            assert after['skeleton']['calendar']['weeks']==4
            assert not after['materials'] and not after['approval']
            assert after['history'][-1]['skeleton']==before['skeleton']
            assert after['history'][-1]['requirements']==before['requirements']
            assert not after['report']['checks'][0]['errors']
