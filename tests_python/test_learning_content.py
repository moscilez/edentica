from io import BytesIO
from docx import Document
from pptx import Presentation
from backend.learning_content import blocks,add_docx
from backend.slides import program_pptx


def test_readable_blocks_and_docx_preserve_headings_code_and_lists():
    text='## Пример\n\nОбъяснение.\n\n- Первый шаг\n- Второй шаг\n\n```python\nprint("<test>")\n    value = 2\n```\n\n> Проверьте результат'
    result=blocks(text)
    assert [x[0] for x in result]==['heading','paragraph','bullet','bullet','code','callout']
    assert '    value = 2' in result[4][1]
    doc=Document();add_docx(doc,text)
    assert doc.paragraphs[0].style.name=='Heading 3'
    assert doc.paragraphs[2].style.name=='List Bullet'
    assert '<test>' in doc.paragraphs[4].text


def test_slide_layouts_are_distinct_and_keep_all_content():
    slides=[{'layout':layout,'title':layout,'text':'- Первый шаг\n- Второй шаг','notes':'Пояснение'} for layout in ['concept','steps','comparison','example','summary']]
    p={'skeleton':{'title':'Курс'},'materials':{'m':{'lectures':[{'slides':slides}]}}}
    deck=Presentation(BytesIO(program_pptx(p)))
    assert len(deck.slides)==6
    for slide in list(deck.slides)[1:]:
        text=' '.join(s.text for s in slide.shapes if s.has_text_frame)
        assert 'Первый шаг' in text and 'Второй шаг' in text
    assert len(deck.slides[2].shapes)!=len(deck.slides[1].shapes)


def test_pedagogical_gate_rejects_plans_and_duplicate_lessons():
    from backend.learning_content import lecture_issues
    from tests_python.lesson_samples import educational_text
    for n in range(1,4):
        lesson={'title':'Учебный урок','topicId':f'M{n}-T1','content':educational_text({'id':f'M{n}-T1'})}
        assert not lecture_issues(lesson),lecture_issues(lesson)
    plan={'topicId':'A','content':('На этом уроке мы будем изучать тему. Вы научитесь решать задачи. Мы рассмотрим примеры. '*70)}
    issues=lecture_issues(plan)
    assert any('обещания' in x for x in issues) and any('пример' in x for x in issues)
    a={'title':'Первый урок','topicId':'A','content':educational_text({'id':'M1-T1'})}
    b={**a,'topicId':'B'}
    assert any('повторяет' in x for x in lecture_issues(b,[a]))
    assert not lecture_issues(a,[a])


def test_structured_blocks_keep_python_indentation_and_reject_flat_code():
    import pytest
    from backend.learning_content import normalize_lecture,lecture_issues
    from tests_python.lesson_samples import educational_text
    from backend.learning_content import blocks
    material={'id':'l','topicId':'T','competencyIds':['C1'],'title':'Ввод','plan':'Практика','slides':[{'title':'Пример','text':'Пример','notes':'Объяснение'}]}
    structured=[]
    for kind,text in blocks(educational_text({'id':'M1-T1'})):
        structured.append({'kind':'paragraph' if kind=='paragraph' else kind,'text':'python' if kind=='code' else text,'lines':text.split('\n') if kind=='code' else []})
    structured.append({'kind':'code','text':'python','lines':['try:','    value = int("2")','except ValueError:','    value = 0']})
    result=normalize_lecture({**material,'blocks':structured})
    assert '\n    value = int' in result['content']
    assert not lecture_issues(result)
    structured[-1]['lines']=['try:\n    value = int("2")\nexcept ValueError:\n    value = 0']
    assert '\n    value = int' in normalize_lecture({**material,'blocks':structured})['content']
    structured[-1]['lines']=['try: value = int("2") except ValueError: value = 0']
    with pytest.raises(ValueError,match='синтаксис'):normalize_lecture({**material,'blocks':structured})


def test_explicit_lesson_sections_are_assembled_into_teaching_material():
    from backend.learning_content import normalize_lecture,lecture_issues
    from tests_python.lesson_samples import educational_text
    import re
    sample=educational_text({'id':'M1-T1'})
    explanation=sample.split('## Разобранный пример')[0].replace('## Объяснение понятий','').strip()
    lecture={'id':'L1','topicId':'T1','competencyIds':['C1'],'title':'Переменные','plan':'План для автора',
        'explanation':[explanation[:450],explanation[450:]],
        'workedExample':{'context':'Исходные данные: цена билета составляет 120 рублей. Покупатель приобретает три билета. Требуется рассчитать полную стоимость покупки.',
          'steps':['Задайте числовые значения price=120 и count=3.','Умножьте цену на количество и выведите произведение.'],
          'code':[{'language':'python','lines':['price = 120','count = 3','print(price * count)']}],
          'result':'Результат равен 360 рублям. Проверка вручную: сумма трёх одинаковых цен 120 + 120 + 120 также равна 360.'},
        'practice':{'task':'Измените цену билета на 150 рублей, а количество на четыре. Сначала рассчитайте стоимость вручную, затем выполните программу и сравните результаты.',
                    'expectedResult':'Программа выводит число 600, соответствующее стоимости четырёх билетов.'},
        'mistakes':[{'error':'Использование строки вместо числа.','correction':'Преобразуйте строку через int до умножения; иначе операция со строкой может повторить символы.'}],
        'selfCheck':[{'question':'Какой тип возвращает функция input?','answer':'Она возвращает строку, поэтому перед вычислением целого количества требуется преобразование int.'},
                     {'question':'Как проверить результат умножения?','answer':'Сложите цену столько раз, сколько приобретено билетов, и сравните ручной результат с выводом.'}],
        'takeaways':['Отделяйте исходные данные от вычисления результата.','Выбирайте числовые типы для выполнения арифметических действий.'],
        'slides':[{'title':'Вычисление стоимости','text':'120 × 3 = 360','notes':'Проверьте стоимость вручную.'}]}
    result=normalize_lecture(lecture)
    assert not lecture_issues(result),lecture_issues(result)
    assert '## Объяснение понятий' in result['content']
    assert '## Самопроверка с ответами' in result['content']
    assert 'Ответ: Она возвращает строку' in result['content']
    assert '```python\nprice = 120' in result['content']
