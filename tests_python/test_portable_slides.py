from io import BytesIO

from fastapi import HTTPException
from pptx import Presentation
import pytest

from backend.slides import program_pptx


def test_portable_export_retains_russian_text_notes_and_splits_long_slides():
    program = {"skeleton": {"title": "Тестовая программа", "audience": "Методисты"},
               "materials": {"m1": {"lectures": [{"slides": [
                   {"title": "Русский заголовок", "text": "Я" * 2100, "notes": "Заметки преподавателя"}]}]}}}
    deck = Presentation(BytesIO(program_pptx(program, "m1")))
    assert len(deck.slides) == 4
    assert "Тестовая программа" in " ".join(s.text for s in deck.slides[0].shapes if s.has_text_frame)
    assert deck.slides[1].notes_slide.notes_text_frame.text == "Заметки преподавателя"
    assert sum(s.text.count("Я") for slide in deck.slides for s in slide.shapes if s.has_text_frame) == 2100


def test_portable_export_rejects_missing_and_empty_modules():
    program = {"skeleton": {"title": "Тест"}, "materials": {}}
    with pytest.raises(HTTPException) as error:
        program_pptx(program, "missing")
    assert error.value.status_code == 409
    with pytest.raises(HTTPException) as error:
        program_pptx(program)
    assert error.value.status_code == 422
