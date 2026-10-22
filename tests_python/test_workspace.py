import asyncio
import copy
import time
from uuid import uuid4
from fastapi.testclient import TestClient
from cryptography.fernet import Fernet
import pytest

from backend.app import create_app
from backend.platform import lessons
from backend.net import public_request, public_url

H={"X-Edentica":"1"}


def reg(client,role):
    client.cookies.clear()
    login=role+'-'+uuid4().hex[:14]
    r=client.post('/api/auth/register',headers=H,json={'login':login,'name':'Тестовый '+role,'password':'Disposable-test-only-123','role':role})
    assert r.status_code==200,r.text
    return r.json(),login


def call(c,action,**data):
    return c.post('/api/platform',headers=H,json={'action':action,**data})


def a(c,action,**data):
    return c.post('/api/agent',headers=H,json={'action':action,**data})


def test_two_party_publication_immutable_lessons_and_learning(client):
    owner,login=reg(client,'methodist')
    p=a(client,'demo').json()
    p['demo']=False
    for module in p['materials'].values():
        for lecture in module['lectures']:
            lecture['content']='Полный учебный материал для тестового сценария. '+('Объяснение понятия, практический пример и проверка понимания. '*8)
    from psycopg.types.json import Jsonb
    client.app.state.db.execute('UPDATE program_projects SET document=%s WHERE id=%s',(Jsonb(p),p['id']))
    assert a(client,'approve',id=p['id'],confirm=True,expectedRevision=p['revision']).status_code==200
    assert call(client,'publish',id=p['id'],confirm=True,digest='no').status_code==409
    assert call(client,'request_review',id=p['id'],expert=owner['id']).status_code==422
    owner_cookie=client.cookies.get('session')
    expert,_=reg(client,'expert');expert_cookie=client.cookies.get('session')
    assert a(client,'get',id=p['id']).status_code==403
    assert client.get('/api/usage').status_code==403
    assert call(client,'learning').status_code==403
    client.cookies.set('session',owner_cookie)
    assert call(client,'request_review',id=p['id'],expert=expert['id']).status_code==200
    review=call(client,'review',id=p['id']).json()
    assert call(client,'publish',id=p['id'],confirm=True,digest=review['digest']).status_code==409
    client.cookies.set('session',expert_cookie)
    assert call(client,'comment',id=p['id'],scope='m1',message='Уточнить результат обучения').status_code==200
    review=call(client,'review',id=p['id']).json()
    assert call(client,'expert_decision',id=p['id'],status='approved',digest=review['digest']).status_code==422
    client.cookies.set('session',owner_cookie)
    assert call(client,'resolve_comment',id=p['id'],commentId=review['comments'][0]['id']).status_code==200
    client.cookies.set('session',expert_cookie)
    assert call(client,'expert_decision',id=p['id'],status='approved',digest=review['digest']).status_code==200
    client.cookies.set('session',owner_cookie)
    assert call(client,'visibility',id=p['id'],visibility='public').status_code==200
    result=call(client,'publish',id=p['id'],confirm=True,digest=review['digest'])
    assert result.status_code==200,result.text
    release=result.json()['id']
    old=call(client,'course',id=release).json()
    material=copy.deepcopy(p['materials'][next(iter(p['materials']))]);material['summary']+=' новая версия'
    r=a(client,'edit_module',id=p['id'],moduleId=material['moduleId'],material=material,expectedRevision=p['revision'])
    assert r.status_code==200,r.text
    changed=call(client,'review',id=p['id']).json()
    assert not changed['current']
    assert call(client,'publish',id=p['id'],confirm=True,digest=changed['digest']).status_code==409
    assert call(client,'course',id=release).json()['description']==old['description']
    learner,learner_login=reg(client,'learner')
    assert a(client,'create',request='Hidden').status_code==403
    assert call(client,'review',id=p['id']).status_code==404
    assert client.get('/api/usage').status_code==403
    assert call(client,'enroll',id=release).status_code==200
    enrolled=call(client,'learning').json()[0]
    units=lessons(p)
    unit=units[0]
    response=call(client,'lesson',enrollment=enrolled['id'],lessonId=unit['id'])
    assert response.status_code==200,response.text
    for q in response.json()['lesson']['questions']:
        assert 'answer' not in q and 'explanation' not in q
    answers={q['id']:q['answer'] for q in unit['questions']}
    finished=call(client,'complete_lesson',enrollment=enrolled['id'],lessonId=unit['id'],answers=answers)
    assert finished.status_code==200 and finished.json()['completed']
    assert call(client,'rate',id=release,rating=5,comment='Понятный материал').status_code==200
    assert call(client,'create_group',title='No').status_code==403
    learner_cookie=client.cookies.get('session')
    client.cookies.set('session',owner_cookie)
    own_stats=call(client,'course_statistics').json()
    assert len(own_stats)==1 and own_stats[0]['completed']==1
    assert own_stats[0]['enrollment_type']=='self' and own_stats[0]['group_title'] is None
    client.cookies.set('session',expert_cookie)
    assert call(client,'enroll',id=release).status_code==403
    assert call(client,'course_statistics').status_code==403
    client.cookies.set('session',owner_cookie)
    group=call(client,'create_group',title='Первая группа').json()['id']
    assert call(client,'add_member',groupId=group,login=learner_login).status_code==200
    assert call(client,'assign',groupId=group,publicationId=release,deadline='2099-01-01T00:00:00Z').status_code==200
    stats=call(client,'statistics',groupId=group)
    assert stats.status_code==200 and stats.json()[0]['completed']==1
    client.cookies.set('session',learner_cookie)
    assert call(client,'statistics',groupId=group).status_code==403
    client.cookies.set('session',owner_cookie)
    assert call(client,'visibility',id=p['id'],visibility='private').status_code==200
    reg(client,'learner')
    assert call(client,'course',id=release).status_code==404
    assert call(client,'catalog').json()==[]
    client.cookies.set('session',learner_cookie)
    assert call(client,'course',id=release).status_code==200


def test_encrypted_models_are_per_user_and_profile_role_cannot_be_escalated(client,monkeypatch):
    monkeypatch.setenv('APP_ENCRYPTION_KEY',Fernet.generate_key().decode())
    owner,_=reg(client,'methodist')
    fake='non-live-unit-test-secret'
    r=client.post('/api/connections',headers=H,json={'action':'save','name':'test','baseUrl':'https://example.com/v1','model':'unit','apiKey':fake})
    assert r.status_code==200,r.text
    assert fake not in r.text and 'encrypted_key' not in r.text
    saved=r.json()[0]['id']
    row=client.app.state.db.one('SELECT encrypted_key FROM model_connections WHERE id=%s',(saved,))
    assert fake not in row['encrypted_key']
    other,_=reg(client,'methodist')
    assert client.post('/api/connections',headers=H,json={}).json()==[]
    assert client.post('/api/connections',headers=H,json={'action':'select','id':saved}).status_code==404
    r=client.post('/api/profile',headers=H,json={'name':'Изменённое ФИО','login':'new-'+uuid4().hex,'role':'expert'})
    assert r.status_code==200 and r.json()['role']=='methodist'
    reg(client,'learner')
    assert client.post('/api/connections',headers=H,json={}).status_code==403


@pytest.mark.parametrize('url',['http://example.com','https://localhost/v1','https://a.local/v1','https://name:secret@example.com','https://example.com:9000/v1'])
def test_llm_endpoint_validation(url):
    with pytest.raises(ValueError):public_url(url)


def test_dns_private_ranges_and_rebinding_cannot_reach_http(monkeypatch):
    monkeypatch.setattr('backend.net.socket.getaddrinfo',lambda *a,**k:[(2,1,6,'',('127.0.0.1',443))])
    with pytest.raises(ValueError,match='служебным'):
        asyncio.run(public_request('GET','https://example.com/'))


def test_lesson_exports_are_portable_and_reimportable():
    from backend.rules import Rules
    from backend.lesson_exports import export_lesson
    from io import BytesIO
    from docx import Document
    from pptx import Presentation
    from pypdf import PdfReader
    p=asyncio.run(Rules().call('demonstration'))
    lesson=lessons(p)[0]
    docx=export_lesson(p,lesson['id'],'docx')
    assert lesson['title'] in '\n'.join(x.text for x in Document(BytesIO(docx.body)).paragraphs)
    assert len(Presentation(BytesIO(export_lesson(p,lesson['id'],'pptx').body)).slides)>1
    assert len(PdfReader(BytesIO(export_lesson(p,lesson['id'],'pdf').body)).pages)>0
    assert len(PdfReader(BytesIO(export_lesson(p,lesson['id'],'slides-pdf').body)).pages)==len(lesson['slides'])


def test_public_http_startup_is_rejected_before_database_access(monkeypatch):
    monkeypatch.setenv('PUBLIC_URL','http://example.com:8000')
    monkeypatch.setenv('COOKIE_SECURE','0')
    with pytest.raises(RuntimeError,match='HTTPS'):
        with TestClient(create_app('postgresql://unreachable:never@127.0.0.1:1/no_db')):
            pass


def test_password_change_requires_current_password_and_revokes_sessions(client):
    account,login=reg(client,'methodist')
    bad=client.post('/api/profile',headers=H,json={'action':'password','currentPassword':'wrong','newPassword':'New-test-password-123'})
    assert bad.status_code==403
    good=client.post('/api/profile',headers=H,json={'action':'password','currentPassword':'Disposable-test-only-123','newPassword':'New-test-password-123'})
    assert good.status_code==200
    assert client.get('/api/me').json()['user'] is None
    assert client.post('/api/auth/login',headers=H,json={'login':login,'password':'Disposable-test-only-123'}).status_code==401
    assert client.post('/api/auth/login',headers=H,json={'login':login,'password':'New-test-password-123'}).status_code==200


def test_public_transport_pins_dns_and_rejects_compression(monkeypatch):
    import httpx
    monkeypatch.setattr('backend.net.socket.getaddrinfo',lambda *a,**k:[(2,1,6,'',('93.184.216.34',443))])
    real=httpx.AsyncClient
    observed=[]
    def handler(request):
        observed.append(request)
        return httpx.Response(200,headers={'content-encoding':'gzip'},stream=httpx.ByteStream(b'compressed payload rejected without decoding'))
    monkeypatch.setattr('backend.net.httpx.AsyncClient',lambda **kwargs:real(transport=httpx.MockTransport(handler),**kwargs))
    with pytest.raises(ValueError,match='Сжатые'):
        asyncio.run(public_request('GET','https://example.com/content'))
    assert observed[0].url.host=='93.184.216.34'
    assert observed[0].headers['host']=='example.com'
    assert observed[0].extensions['sni_hostname']=='example.com'


def test_complete_brief_builds_outline_in_one_call_without_approving(postgres_url):
    from backend.rules import Rules
    s=asyncio.run(Rules().call('demonstration'))['skeleton']
    calls=[]
    async def complete(skill,payload):
        calls.append(payload)
        await asyncio.sleep(.1)
        return {'value':{'requirements':{k:s[k] for k in ('type','totalHours','audience','entryLevel','format','assessment')},'questions':[],'notes':[],'skeleton':s},'usage':{'inputTokens':100,'outputTokens':500},'model':'test-model'}
    with TestClient(create_app(postgres_url,complete),base_url='http://localhost') as client:
        reg(client,'methodist')
        p=a(client,'create',request='Общеразвивающая программа на 16 часов').json()
        second=a(client,'create',request='Общеразвивающая программа на 16 часов').json()
        a(client,'analyze',id=p['id'])
        a(client,'analyze',id=second['id'])
        deadline=time.monotonic()+15
        while time.monotonic()<deadline:
            result=a(client,'get',id=p['id']).json()
            other=a(client,'get',id=second['id']).json()
            if result.get('job',{}).get('state')!='running' and other.get('job',{}).get('state')!='running':break
            time.sleep(.05)
        assert result['state']=='skeleton_review',result
        assert len(calls)==1
        assert result['usage'].get('cacheHits',0)+other['usage'].get('cacheHits',0)==1
        assert result['approval'] is None
        assert a(client,'generate',id=p['id']).status_code==422


def test_public_signup_cannot_spend_server_model_quota(monkeypatch):
    from backend.models import Models
    from fastapi import HTTPException
    monkeypatch.setenv('EDENTICA_MODE','hackathon')
    monkeypatch.setenv('HOST','0.0.0.0')
    monkeypatch.setenv('SYSTEM_LLM_SHARED_ACCESS','0')
    monkeypatch.setenv('SYSTEM_LLM_ALLOWED_USERS','approved-immutable-id')
    models=Models(None)
    assert not asyncio.run(models.system_allowed('new-user'))
    assert asyncio.run(models.system_allowed('approved-immutable-id'))
    async def forbidden(*args):raise AssertionError('Server model must not be called')
    with pytest.raises(HTTPException) as error:
        asyncio.run(models.complete('new-user','instruction',{},forbidden,selected=None))
    assert error.value.status_code==403


def test_only_admin_transport_can_allow_private_gateway(monkeypatch):
    import httpx
    from backend.net import public_request
    monkeypatch.setattr('backend.net.socket.getaddrinfo',lambda *a,**k:[(2,1,6,'',('10.250.0.2',443))])
    real=httpx.AsyncClient
    monkeypatch.setattr('backend.net.httpx.AsyncClient',lambda **kwargs:real(transport=httpx.MockTransport(lambda r:httpx.Response(200,stream=httpx.ByteStream(b'{}'))),**kwargs))
    with pytest.raises(ValueError,match='служебным'):
        asyncio.run(public_request('GET','https://gateway.example/v1'))
    assert asyncio.run(public_request('GET','https://gateway.example/v1',allow_private=True))[0]==200
    monkeypatch.setattr('backend.net.socket.getaddrinfo',lambda *a,**k:[(2,1,6,'',('169.254.169.254',443))])
    with pytest.raises(ValueError,match='служебным'):
        asyncio.run(public_request('GET','https://gateway.example/v1',allow_private=True))


def test_malformed_intake_is_repaired_not_cached(postgres_url):
    from backend.rules import Rules
    s=asyncio.run(Rules().call('demonstration'))['skeleton']
    calls=[]
    async def complete(skill,payload):
        calls.append(payload)
        value={'requirements':{k:s[k] for k in ('type','totalHours','audience','entryLevel','format','assessment')},'questions':[],'notes':[],'skeleton':s}
        if len(calls)==1:value['questions']=['wrong shape']
        return {'value':value,'usage':{'inputTokens':10,'outputTokens':20},'model':'test-model'}
    with TestClient(create_app(postgres_url,complete),base_url='http://localhost') as client:
        reg(client,'methodist')
        p=a(client,'create',request='Общеразвивающая программа на 16 часов. Проверка восстановления схемы '+uuid4().hex).json()
        a(client,'analyze',id=p['id'])
        deadline=time.monotonic()+15
        while time.monotonic()<deadline:
            result=a(client,'get',id=p['id']).json()
            if result.get('job',{}).get('state')!='running':break
            time.sleep(.05)
        assert result['job']['state']=='completed',result['job']
        assert result['state']=='skeleton_review'
        assert len(calls)==2 and calls[1]['validationErrors']
        assert result['usage']['calls']==2
        assert client.app.state.db.one('SELECT count(*) AS n FROM llm_cache WHERE owner=%s',(p['owner'],))['n']==1


def test_complete_brief_does_not_ask_to_repeat_extracted_audience(postgres_url):
    from backend.rules import Rules
    s=asyncio.run(Rules().call('demonstration'))['skeleton']
    async def complete(skill,payload):
        return {'value':{'requirements':{k:s[k] for k in ('type','totalHours','audience','entryLevel','format','assessment')},'questions':[{'field':'аудитория','question':'Для школьников какого возраста?'}],'notes':[],'skeleton':s},'usage':{},'model':'test-model'}
    with TestClient(create_app(postgres_url,complete),base_url='http://localhost') as client:
        reg(client,'methodist')
        p=a(client,'create',request='Общеразвивающая программа на 16 часов со всеми обязательными полями').json()
        a(client,'analyze',id=p['id'])
        deadline=time.monotonic()+15
        while time.monotonic()<deadline:
            result=a(client,'get',id=p['id']).json()
            if result.get('job',{}).get('state')!='running':break
            time.sleep(.05)
        assert result['state']=='skeleton_review',result
        assert result['questions']==[]


def test_truncated_combined_outline_falls_back_to_intake(postgres_url):
    from fastapi import HTTPException
    from backend.rules import Rules
    s=asyncio.run(Rules().call('demonstration'))['skeleton']
    calls=[]
    async def complete(skill,payload):
        calls.append(skill)
        if len(calls)==1:raise HTTPException(422,'Ответ обрезан. Увеличьте лимит.')
        return {'value':{'requirements':{k:s[k] for k in ('type','totalHours','audience','entryLevel','format','assessment')},'questions':[],'notes':[]},'usage':{},'model':'test-model'}
    with TestClient(create_app(postgres_url,complete),base_url='http://localhost') as client:
        reg(client,'methodist')
        p=a(client,'create',request='Общеразвивающая программа на 16 часов. Проверка запасного пути '+uuid4().hex).json()
        a(client,'analyze',id=p['id'])
        deadline=time.monotonic()+15
        while time.monotonic()<deadline:
            result=a(client,'get',id=p['id']).json()
            if result.get('job',{}).get('state')!='running':break
            time.sleep(.05)
        assert result['job']['state']=='completed',result['job']
        assert result['state']=='ready_for_skeleton'
        assert len(calls)==2
        assert all('Навык 2' not in skill for skill in calls)  # Intake stays compact on both attempts


def test_review_waits_for_complete_course_and_module_generation_is_scoped(client,monkeypatch):
    owner,_=reg(client,'methodist');owner_cookie=client.cookies.get('session')
    expert,_=reg(client,'expert');expert_cookie=client.cookies.get('session');client.cookies.set('session',owner_cookie)
    p=a(client,'demo').json();p['demo']=False;p['materials']={};p['sections']={};p['finalAssessment']=None
    from psycopg.types.json import Jsonb
    client.app.state.db.execute('UPDATE program_projects SET document=%s WHERE id=%s',(Jsonb(p),p['id']))
    assert a(client,'approve',id=p['id'],confirm=True,expectedRevision=p['revision']).status_code==200
    result=call(client,'request_review',id=p['id'],expert=expert['id'])
    assert result.status_code==409 and 'завершите' in result.text
    assert not client.app.state.db.one('SELECT 1 FROM program_reviews WHERE project_id=%s',(p['id'],))
    # Old pending reviews cannot expose an unfinished revision to the expert.
    client.app.state.db.execute('INSERT INTO program_reviews(project_id,expert,digest) VALUES(%s,%s,%s)',(p['id'],expert['id'],'old-digest'))
    client.cookies.set('session',expert_cookie)
    assert call(client,'review',id=p['id']).status_code==409
    assert call(client,'comment',id=p['id'],message='Premature comment').status_code==409
    client.cookies.set('session',owner_cookie)
    targets=[]
    async def generate(current,target=None,instruction=None):
        targets.append(target['id']);current['state']='approved'
    monkeypatch.setattr(client.app.state.programs,'generate',generate)
    mid=p['skeleton']['modules'][0]['id']
    assert a(client,'generate_module',id=p['id'],moduleId='missing').status_code==404
    assert a(client,'generate_module',id=p['id'],moduleId=mid).status_code==200
    deadline=time.monotonic()+5
    while time.monotonic()<deadline:
        current=a(client,'get',id=p['id']).json()
        if current['job']['state']!='running':break
        time.sleep(.02)
    assert targets==[mid]
    assert current['materials']=={} and current['sections']=={}
